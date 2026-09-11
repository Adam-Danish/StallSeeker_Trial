from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(__file__).with_name("Task Checklist Current.docx")
OUTPUT = Path(__file__).with_name("Task Checklist Simplified.docx")


SECTIONS = [
    ("INFORMATION GATHERING", [
        ("1", "Finalise the problem statement, objectives, scope, and requirements", "TBC", "All members", ""),
        ("2", "Complete the system architecture, database, interface, and use-case designs", "TBC", "Adam Danish", ""),
    ]),
    ("DEVELOPMENT", [
        ("3", "Set up the Flutter project, dependencies, and application theme", "27/07/2026", "Adam Danish", "X"),
        ("4", "Configure Firebase and create the application models and service layer", "27/07/2026-05/08/2026", "Adam Danish", "X"),
        ("5", "Implement authentication, user roles, guest access, and account management", "29/07/2026-06/09/2026", "Adam Danish", "X"),
        ("6", "Build vendor profiles, stall information, images, and opening hours", "03/08/2026-14/08/2026", "Adam Danish", "X"),
        ("7", "Build the vendor dashboard, location sharing, and stall status", "03/08/2026-25/08/2026", "Adam Danish", "X"),
        ("8", "Build vendor menu creation, editing, deletion, and availability", "03/08/2026-14/08/2026", "Adam Danish", "X"),
        ("9", "Build customer map discovery using GPS and manual locations", "03/08/2026-06/09/2026", "Adam Danish", "X"),
        ("10", "Add stall search, filters, details, menus, and directions", "03/08/2026-06/09/2026", "Adam Danish", "X"),
        ("11", "Implement following, push notifications, history, and preferences", "05/08/2026-06/09/2026", "Adam Danish", "X"),
        ("12", "Complete profile, settings, navigation, and interface improvements", "06/08/2026-06/09/2026", "Adam Danish", "X"),
        ("13", "Configure production security rules, indexes, and Cloud Functions", "TBC", "Adam Danish", ""),
        ("14", "Prepare the signed Android release and installation package", "TBC", "Adam Danish", ""),
    ]),
    ("TESTING", [
        ("15", "Run backend tests and Flutter static analysis", "10/09/2026-TBC", "Adam Danish", ""),
        ("16", "Test authentication, vendor, customer, map, and notification functions", "TBC", "Adam Danish", ""),
        ("17", "Conduct security, device, usability, and release-build testing", "TBC", "Adam Danish", ""),
    ]),
    ("DOCUMENTATION", [
        ("18", "Update the report chapters for design, implementation, and testing", "TBC", "Adam Danish", ""),
        ("19", "Prepare the user manual, deployment guide, diagrams, and screenshots", "TBC", "Adam Danish", ""),
        ("20", "Complete references, appendices, proofreading, and final formatting", "TBC", "Adam Danish", ""),
    ]),
    ("PROGRESS PRESENTATION #1", [
        ("21", "Prepare and deliver Progress Presentation 1 with a system demo", "TBC", "Adam Danish", ""),
    ]),
    ("PROGRESS PRESENTATION #2", [
        ("22", "Prepare and deliver Progress Presentation 2 with testing evidence", "TBC", "Adam Danish", ""),
    ]),
    ("INTERNAL PRESENTATION", [
        ("23", "Deliver the internal presentation and apply panel corrections", "TBC", "Adam Danish", ""),
    ]),
    ("EXTERNAL PRESENTATION", [
        ("24", "Finalise all materials and deliver the external presentation", "TBC", "Adam Danish", ""),
    ]),
]


def replace_cell_text(cell, value):
    first = cell.paragraphs[0]
    for paragraph in cell.paragraphs[1:]:
        cell._tc.remove(paragraph._p)
    if first.runs:
        first.runs[0].text = value
        for run in first.runs[1:]:
            first._p.remove(run._r)
    else:
        first.add_run(value)


def prepare_row(row, repeat_header=False):
    tr_pr = row._tr.get_or_add_trPr()
    for height in list(tr_pr.findall(qn("w:trHeight"))):
        tr_pr.remove(height)
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))
    if repeat_header:
        header = tr_pr.find(qn("w:tblHeader"))
        if header is None:
            header = OxmlElement("w:tblHeader")
            tr_pr.append(header)
        header.set(qn("w:val"), "1")


doc = Document(SOURCE)
table = doc.tables[1]

header_template = deepcopy(table.rows[0]._tr)
section_template = deepcopy(table.rows[1]._tr)
task_template = deepcopy(table.rows[2]._tr)

for row in list(table.rows):
    table._tbl.remove(row._tr)

table._tbl.append(deepcopy(header_template))
prepare_row(table.rows[-1], repeat_header=True)

for section_name, tasks in SECTIONS:
    table._tbl.append(deepcopy(section_template))
    section_row = table.rows[-1]
    for cell, value in zip(section_row.cells, ("", section_name, "", "", "")):
        replace_cell_text(cell, value)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    prepare_row(section_row)

    for task in tasks:
        table._tbl.append(deepcopy(task_template))
        row = table.rows[-1]
        for cell, value in zip(row.cells, task):
            replace_cell_text(cell, value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        prepare_row(row)

# Remove the continuation table created by the longer version and its empty
# page-break paragraph, leaving the user's title, metadata, and instruction.
if len(doc.tables) > 2:
    continuation = doc.tables[2]._tbl
    previous = continuation.getprevious()
    continuation.getparent().remove(continuation)
    if previous is not None and previous.tag == qn("w:p") and not "".join(previous.itertext()).strip():
        previous.getparent().remove(previous)

doc.save(OUTPUT)
print(OUTPUT)
