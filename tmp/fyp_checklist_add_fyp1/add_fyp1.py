from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(__file__).with_name("Task Checklist Current.docx")
OUTPUT = Path(__file__).with_name("Task Checklist Updated.docx")


FYP1_BLOCK = [
    ("section", "FYP 01"),
    ("section", "PLANNING AND ANALYSIS PHASE"),
    ("task", ("1", "Individual proposal brainstorming", "", "", "")),
    ("task", ("2", "Individual proposal presentation and project selection", "", "", "")),
    ("section", "PROJECT DESIGN AND CONCEPT DEVELOPMENT"),
    ("task", ("3", "Conduct feasibility study and prepare flowchart", "", "", "")),
    ("task", ("4", "Prepare use-case, activity, sequence, and class diagrams", "", "", "")),
    ("task", ("5", "Create GUI design prototype", "", "", "")),
    ("section", "FYP 2"),
]


def replace_cell_text(cell, value):
    first = cell.paragraphs[0]
    for paragraph in cell.paragraphs[1:]:
        cell._tc.remove(paragraph._p)
    if first.runs:
        first.runs[0].text = value
        for run in first.runs[1:]:
            first._p.remove(run._r)
    else:
        first.add_run(value)


def prepare_row(row, repeat_header=False):
    tr_pr = row._tr.get_or_add_trPr()
    for height in list(tr_pr.findall(qn("w:trHeight"))):
        tr_pr.remove(height)
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))
    if repeat_header:
        header = tr_pr.find(qn("w:tblHeader"))
        if header is None:
            header = OxmlElement("w:tblHeader")
            tr_pr.append(header)
        header.set(qn("w:val"), "1")


doc = Document(SOURCE)
table = doc.tables[1]

header_template = deepcopy(table.rows[0]._tr)
section_template = deepcopy(table.rows[1]._tr)
task_template = next(
    deepcopy(row._tr)
    for row in table.rows[1:]
    if row.cells[0].text.strip().isdigit()
)
existing_rows = [[cell.text for cell in row.cells] for row in table.rows[1:]]

for row in list(table.rows):
    table._tbl.remove(row._tr)

table._tbl.append(header_template)
prepare_row(table.rows[-1], repeat_header=True)

for kind, content in FYP1_BLOCK:
    if kind == "section":
        table._tbl.append(deepcopy(section_template))
        row = table.rows[-1]
        for cell, value in zip(row.cells, ("", content, "", "", "")):
            replace_cell_text(cell, value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        prepare_row(row)
    else:
        table._tbl.append(deepcopy(task_template))
        row = table.rows[-1]
        for cell, value in zip(row.cells, content):
            replace_cell_text(cell, value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        prepare_row(row)

for values in existing_rows:
    is_task = values[0].strip().isdigit()
    table._tbl.append(deepcopy(task_template if is_task else section_template))
    row = table.rows[-1]
    if is_task:
        values[0] = str(int(values[0].strip()) + 5)
    for cell, value in zip(row.cells, values):
        replace_cell_text(cell, value)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if is_task:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    prepare_row(row)

doc.save(OUTPUT)
print(OUTPUT)
