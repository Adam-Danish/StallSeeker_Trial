from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document

path = Path(r"C:\GMi\SEM 5\FYP 2\Task Checklist Template (1).docx")
doc = Document(path)

print("PARAGRAPHS")
for i, paragraph in enumerate(doc.paragraphs):
    if paragraph.text.strip():
        print(f"P{i}: style={paragraph.style.name!r} text={paragraph.text!r}")

print("\nTABLES")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti}: rows={len(table.rows)} cols={len(table.columns)} style={table.style.name if table.style else None!r}")
    for ri, row in enumerate(table.rows):
        values = [" | ".join(p.text for p in cell.paragraphs) for cell in row.cells]
        print(f"  R{ri}: {values!r}")

print("\nSECTIONS")
for si, section in enumerate(doc.sections):
    print(si, section.page_width, section.page_height, section.top_margin, section.bottom_margin, section.left_margin, section.right_margin)
    print("  HEADER", [p.text for p in section.header.paragraphs])
    print("  FOOTER", [p.text for p in section.footer.paragraphs])

print("\nPACKAGE")
with ZipFile(path) as archive:
    for info in archive.infolist():
        print(f"{info.filename}\t{info.file_size}\t{info.CRC}")

