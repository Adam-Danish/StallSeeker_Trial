from pathlib import Path
from zipfile import ZipFile
import hashlib
from docx import Document
from docx.oxml.ns import qn

path = Path(r"C:\GMi\SEM 5\FYP 2\Task Checklist Template (1).docx")
doc = Document(path)

for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti} autofit={table.autofit}")
    print("  grid widths", [c.width for c in table.columns])
    for ri, row in enumerate(table.rows):
        print(f"  ROW {ri} height={row.height} rule={row.height_rule}")
        for ci, cell in enumerate(row.cells):
            tcPr = cell._tc.tcPr
            shd = tcPr.find(qn('w:shd'))
            fill = shd.get(qn('w:fill')) if shd is not None else None
            tc_mar = tcPr.find(qn('w:tcMar'))
            margins = {}
            if tc_mar is not None:
                for side in ('top', 'left', 'bottom', 'right', 'start', 'end'):
                    element = tc_mar.find(qn(f'w:{side}'))
                    if element is not None:
                        margins[side] = element.get(qn('w:w'))
            pinfo = []
            for p in cell.paragraphs:
                runs = []
                for run in p.runs:
                    runs.append({
                        'text': run.text,
                        'font': run.font.name,
                        'size': run.font.size.pt if run.font.size else None,
                        'bold': run.bold,
                        'underline': run.underline,
                    })
                pinfo.append({'text': p.text, 'alignment': p.alignment, 'style': p.style.name, 'runs': runs})
            print(f"    C{ci} width={cell.width} valign={cell.vertical_alignment} fill={fill} margins={margins} p={pinfo}")

print("PACKAGE_HASHES")
with ZipFile(path) as archive:
    for info in archive.infolist():
        data = archive.read(info.filename)
        print(f"{info.filename}\t{len(data)}\t{hashlib.sha256(data).hexdigest()}")
