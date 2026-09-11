from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(__file__).with_name("Task Checklist Template (1).docx")
OUTPUT = Path(__file__).with_name("Task Checklist Template (1) - Completed.docx")


SECTIONS = [
    ("INFORMATION GATHERING", [
        ("1", "Confirm the problem statement, project objectives, scope, and success criteria", "TBC", "Adam Danish", ""),
        ("2", "Finalise requirements and user stories for customer, vendor, and guest users", "TBC", "Adam Danish", ""),
        ("3", "Complete the system architecture, database design, interface design, and use-case diagrams", "TBC", "Adam Danish", ""),
    ]),
    ("DEVELOPMENT", [
        ("4", "Set up the Flutter project structure, dependencies, and application theme", "27/07/2026", "Adam Danish", "X"),
        ("5", "Configure Firebase Core, Authentication, Firestore, and Firebase Storage", "27/07/2026-04/08/2026", "Adam Danish", "X"),
        ("6", "Create the application models, Firestore collection constants, and service layer", "27/07/2026-05/08/2026", "Adam Danish", "X"),
        ("7", "Implement email and password registration and login", "29/07/2026-06/09/2026", "Adam Danish", "X"),
        ("8", "Implement customer and vendor role selection and screen routing", "29/07/2026-06/09/2026", "Adam Danish", "X"),
        ("9", "Implement guest access and Google Sign-In", "06/08/2026-12/08/2026", "Adam Danish", "X"),
        ("10", "Implement email verification, forgotten password, change password, and change email functions", "06/08/2026-06/09/2026", "Adam Danish", "X"),
        ("11", "Develop the vendor profile and stall information management screens", "03/08/2026-14/08/2026", "Adam Danish", "X"),
        ("12", "Implement stall image upload, description, category, and opening-hours management", "03/08/2026-14/08/2026", "Adam Danish", "X"),
        ("13", "Develop the vendor dashboard with location sharing and open or closed stall status", "03/08/2026-25/08/2026", "Adam Danish", "X"),
        ("14", "Develop vendor menu creation, editing, deletion, pricing, images, and availability", "03/08/2026-14/08/2026", "Adam Danish", "X"),
        ("15", "Develop customer stall discovery using GPS and Google Maps", "03/08/2026-06/09/2026", "Adam Danish", "X"),
        ("16", "Implement manual location selection for customers", "12/08/2026-06/09/2026", "Adam Danish", "X"),
        ("17", "Implement stall search, category filters, distance filters, and map-bound filtering", "03/08/2026-06/09/2026", "Adam Danish", "X"),
        ("18", "Develop stall cards and the stall-details screen with menu information", "03/08/2026-06/09/2026", "Adam Danish", "X"),
        ("19", "Implement stall following and unfollowing with a customer following list", "05/08/2026-25/08/2026", "Adam Danish", "X"),
        ("20", "Implement push notifications for vendor stall-status updates", "06/08/2026-06/09/2026", "Adam Danish", "X"),
        ("21", "Implement notification history, unread status, and notification preferences", "06/09/2026", "Adam Danish", "X"),
        ("22", "Add directions through Google Maps, Waze, and browser fallback", "25/08/2026", "Adam Danish", "X"),
        ("23", "Complete profile, account, About, FAQ, legal, logout, and account-deletion screens", "06/08/2026-06/09/2026", "Adam Danish", "X"),
        ("24", "Improve the splash screen, welcome screen, navigation, loading states, and error messages", "06/08/2026-06/09/2026", "Adam Danish", "X"),
        ("25", "Configure production secrets, Firestore rules, indexes, and Cloud Functions deployment", "TBC", "Adam Danish", ""),
        ("26", "Prepare the signed Android release build and final installation package", "TBC", "Adam Danish", ""),
    ]),
    ("TESTING", [
        ("27", "Run notification-history backend unit tests", "10/09/2026", "Adam Danish", "X"),
        ("28", "Run Flutter static analysis and resolve warnings and errors", "TBC", "Adam Danish", ""),
        ("29", "Test registration, verification, login, Google login, guest access, and role routing", "TBC", "Adam Danish", ""),
        ("30", "Test vendor profile, stall location, operating status, and menu management", "TBC", "Adam Danish", ""),
        ("31", "Test customer location, stall search, filters, details, and directions", "TBC", "Adam Danish", ""),
        ("32", "Test following, notifications, notification history, and notification preferences", "TBC", "Adam Danish", ""),
        ("33", "Validate Firestore, Storage, and Cloud Functions security permissions", "TBC", "Adam Danish", ""),
        ("34", "Conduct Android device, usability, regression, and release-build testing", "TBC", "Adam Danish", ""),
    ]),
    ("DOCUMENTATION", [
        ("35", "Update Chapters 1-3 with the final scope, literature review, methodology, and system design", "TBC", "Adam Danish", ""),
        ("36", "Write the implementation chapter with architecture and feature evidence", "TBC", "Adam Danish", ""),
        ("37", "Write the testing and evaluation chapter with test cases, results, and limitations", "TBC", "Adam Danish", ""),
        ("38", "Prepare the user manual and deployment guide", "TBC", "Adam Danish", ""),
        ("39", "Update diagrams, screenshots, references, appendices, and traceability evidence", "TBC", "Adam Danish", ""),
        ("40", "Proofread and format the final report according to institutional requirements", "TBC", "Adam Danish", ""),
    ]),
    ("PROGRESS PRESENTATION #1", [
        ("41", "Prepare slides, progress evidence, and demonstration flow for Progress Presentation 1", "TBC", "Adam Danish", ""),
        ("42", "Deliver Progress Presentation 1 and record reviewer feedback", "TBC", "Adam Danish", ""),
    ]),
    ("PROGRESS PRESENTATION #2", [
        ("43", "Prepare updated slides, test evidence, and demo for Progress Presentation 2", "TBC", "Adam Danish", ""),
    ]),
    ("INTERNAL PRESENTATION", [
        ("44", "Complete the internal presentation and apply panel corrections", "TBC", "Adam Danish", ""),
    ]),
    ("EXTERNAL PRESENTATION", [
        ("45", "Freeze the tested application, report, user manual, slides, and backup demonstration materials", "TBC", "Adam Danish", ""),
        ("46", "Deliver the external presentation and final system demonstration", "TBC", "Adam Danish", ""),
    ]),
]


def replace_cell_text(cell, value):
    paragraphs = cell.paragraphs
    first = paragraphs[0]
    for paragraph in paragraphs[1:]:
        cell._tc.remove(paragraph._p)
    runs = first.runs
    if runs:
        runs[0].text = value
        for run in runs[1:]:
            first._p.remove(run._r)
    else:
        first.add_run(value)


def remove_fixed_height(row):
    tr_pr = row._tr.get_or_add_trPr()
    for height in list(tr_pr.findall(qn("w:trHeight"))):
        tr_pr.remove(height)
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "1")


doc = Document(SOURCE)
table = doc.tables[1]

header_template = deepcopy(table.rows[0]._tr)
section_template = deepcopy(table.rows[1]._tr)
task_template = deepcopy(table.rows[2]._tr)

for row in list(table.rows):
    table._tbl.remove(row._tr)

table._tbl.append(deepcopy(header_template))
header_row = table.rows[-1]
remove_fixed_height(header_row)
set_repeat_header(header_row)

# Retain the document's table format on a deliberate third-page continuation.
# Splitting at this point prevents a section heading from being stranded at the
# bottom of page 2 and keeps the repeated column headings fully visible.
continuation_xml = deepcopy(table._tbl)


def append_sections(target_table, sections):
    for section_name, tasks in sections:
        target_table._tbl.append(deepcopy(section_template))
        section_row = target_table.rows[-1]
        section_values = ("", section_name, "", "", "")
        for cell, value in zip(section_row.cells, section_values):
            replace_cell_text(cell, value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        remove_fixed_height(section_row)

        for task in tasks:
            target_table._tbl.append(deepcopy(task_template))
            row = target_table.rows[-1]
            for cell, value in zip(row.cells, task):
                replace_cell_text(cell, value)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
            remove_fixed_height(row)


append_sections(table, SECTIONS[:6])

page_break = doc.add_paragraph()
page_break.paragraph_format.page_break_before = True
page_break.paragraph_format.space_before = Pt(0)
page_break.paragraph_format.space_after = Pt(0)
table._tbl.addnext(page_break._p)
page_break._p.addnext(continuation_xml)
continuation_table = doc.tables[-1]
append_sections(continuation_table, SECTIONS[6:])

doc.save(OUTPUT)
print(OUTPUT)
